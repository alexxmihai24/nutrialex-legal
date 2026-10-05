"""Actualiza precios.json con el precio medio en España del Panel de consumo alimentario en hogares (MAPA).

Fuente: https://www.mapa.gob.es/es/alimentacion/temas/consumo-tendencias/panel-de-consumo-alimentario/series-anuales
Fichero usado: el más reciente "AAAAmes-sd-datos-mensuales-panel-consumo-hogares-*.xlsx" (total España, una hoja por mes).
Lo ejecuta .github/workflows/precios.yml cada semana; también se puede lanzar a mano: python scripts/update_prices.py
"""
import io
import json
import os
import re
import ssl
import sys
import urllib.request
from datetime import date

import openpyxl  # con defusedxml instalado, openpyxl lo usa para leer el XML de forma segura

PAGE = "https://www.mapa.gob.es/es/alimentacion/temas/consumo-tendencias/panel-de-consumo-alimentario/series-anuales"
OUT = "precios.json"
# La web del MAPA usa la raiz "AC RAIZ FNMT-RCM SERVIDORES SEGUROS G2R", que Ubuntu aun no trae.
# Copia en certs/ (SHA-1 19:87:94:8D:AB:7D:62:00:9E:2D:6B:BE:D8:83:DC:3A:F5:67:99:E8, la misma que
# Windows tiene en sus raices de confianza). Se anade a las del sistema: la verificacion TLS sigue activa.
TLS = ssl.create_default_context()
TLS.load_verify_locations(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "certs", "fnmt-raiz-servidores-seguros-g2r.pem"))
MONTHS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
FILE_RE = re.compile(r'href="([^"]*?(\d{4})(' + "|".join(MONTHS) + r')-sd-datos-mensuales-panel-consumo-hogares[^"]*?\.xlsx)"', re.I)

# Palabra clave de la app (lib/shoppingList.ts) -> línea del panel. Si una línea sale repetida en la hoja,
# vale la primera: es la de producto fresco (las siguientes son congelado o conserva).
# Lo que no está aquí (tofu, quinoa, avena, hummus...) no tiene línea propia en el panel: la app usa su estimación.
MAPA_LINES = {
    "pollo": "POLLO FILETES",
    "pavo": "PAVO",
    "salmon": "SALMON FRESCO",
    "merluza": "MERLUZA/PESCAD",
    "atun": "CONS.ATUN",
    "huevo": "T.HUEVOS KGS",
    "jamon": "JAMON C.N.LONCHA",
    "brocoli": "BROCOLI",
    "espinacas": "VERDURAS DE HOJA",
    "tomate triturado": "TOMATE N.TRITUR.",
    "tomate": "TOMATES",
    "calabacin": "CALABACINES",
    "zanahoria": "ZANAHORIAS",
    "pimiento": "PIMIENTOS",
    "cebolla": "CEBOLLAS",
    "ajo": "AJOS",
    "limon": "LIMONES",
    "patata": "PATATAS FRESCAS",
    "aguacate": "AGUACATE",
    "platano": "PLATANOS",
    "manzana": "MANZANAS",
    "pepino": "PEPINOS",
    "lechuga": "LECHUGA/ESC./ENDIVIA",
    "esparrago": "ESPARRAGOS",
    "calabaza": "CALABAZA",
    "puerro": "PUERRO",
    "berenjena": "BERENJENAS",
    "arandano": "ARANDANO",
    "champiñon": "CHAMPIÑONES+O.SETAS",
    "arroz": "ARROZ",
    "pasta": "PASTA ALIM.SECA",
    "lentejas": "LENTEJAS SECAS",
    "garbanzos": "GARBANZOS",
    "pan": "PAN",
    "yogur": "YOGUR",
    "leche de almendra": "BEBIDAS VEGETALES",
    "leche": "TOTAL LECHE LIQUIDA",
    "queso fresco": "QUESO FRESCO",
    "queso de cabra": "QUESO LECHE CABRA",
    "parmesano": "QUESO CURADO",
    "queso": "QUESO",
    "aceite": "TOTAL ACEITES OLIVA",
    "nueces": "NUECES SIN CASCARA",
    "almendras": "ALMENDRA SIN CASC.",
    "aceitunas": "ACEITUNAS",
    "miel": "MIEL",
    "salsa": "SALSAS",
}


def fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "NutriAlexAI-precios/1.0 (+https://alexxmihai24.github.io/nutrialex-legal/)"})
    with urllib.request.urlopen(req, timeout=180, context=TLS) as r:
        return r.read()


def latest_file() -> tuple[str, str]:
    html = fetch(PAGE).decode("utf-8", "replace")
    found = [(int(y), MONTHS.index(m.lower()) + 1, href) for href, y, m in FILE_RE.findall(html)]
    if not found:
        sys.exit("No encuentro el Excel mensual en la página del MAPA (¿ha cambiado la web?)")
    y, m, href = max(found)
    url = href if href.startswith("http") else "https://www.mapa.gob.es" + href
    return f"{y}-{m:02d}", url


def read_prices(xlsx: bytes) -> dict[str, float]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[-1]]  # una hoja por mes; la última es el mes más reciente
    rows = ws.iter_rows(values_only=True)
    next(rows)
    header = next(rows)
    col = next((i for i, v in enumerate(header) if isinstance(v, str) and v.upper().startswith("PRECIO MEDIO")), None)
    if col is None:
        sys.exit("No encuentro la columna 'PRECIO MEDIO' (¿ha cambiado el formato del Excel?)")
    prices: dict[str, float] = {}
    for row in rows:
        name, value = row[0], row[col] if len(row) > col else None
        if isinstance(name, str) and isinstance(value, (int, float)) and value > 0:
            prices.setdefault(name.strip(), float(value))  # primera aparición = producto fresco
    return prices


def main() -> None:
    period, url = latest_file()
    try:
        current = json.load(open(OUT, encoding="utf-8"))
    except FileNotFoundError:
        current = {}
    if current.get("period") == period and "--force" not in sys.argv:
        print(f"Sin cambios: {OUT} ya tiene {period}")
        return

    table = read_prices(fetch(url))
    per100g, missing = {}, []
    for key, line in MAPA_LINES.items():
        eur_kg = table.get(line)
        if eur_kg is None:
            missing.append(line)
        elif 0.1 <= eur_kg <= 200:  # €/kg o €/l; fuera de ese rango es un error de lectura
            per100g[key] = round(eur_kg / 10, 3)
    if missing:
        print("Líneas no encontradas (la app usará su estimación):", ", ".join(missing))
    if len(per100g) < len(MAPA_LINES) * 0.8:
        sys.exit(f"Solo {len(per100g)} de {len(MAPA_LINES)} precios: no se actualiza (¿ha cambiado el formato?)")

    out = {
        "source": "MAPA, Panel de consumo alimentario en hogares (precio medio en España)",
        "period": period,
        "updatedAt": date.today().isoformat(),
        "url": url,
        "pricePer100g": per100g,  # euros por 100 g o 100 ml
    }
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print(f"{OUT}: {len(per100g)} precios de {period}")


if __name__ == "__main__":
    main()
